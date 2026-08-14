from __future__ import annotations

import json
import math
import re
import statistics
import wave
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "docs" / "so-sanh-moonshine-v2-va-whisper.docx"
RESULTS_PATH = ROOT / "benchmark" / "results" / "three-model-tts-results.json"
AUDIO_DIR = ROOT / "public" / "benchmark-audio"

INK = "172B3A"
BLUE = "2E74B5"
MUTED = "52606D"
HEADER_FILL = "E7EEF6"
SMALL_FILL = "EAF4FB"
BORDER = "9AA7B2"

MODEL_ORDER = [
    "small-streaming",
    "medium-streaming",
    "whisper-base-en-q5_1",
]
MODEL_NAMES = {
    "small-streaming": "Moonshine Small v2",
    "medium-streaming": "Moonshine Medium v2",
    "whisper-base-en-q5_1": "Whisper Base.en Q5_1",
}


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9\s]", " ", text.lower())).strip()


def word_error_rate(reference: str, hypothesis: str) -> float:
    left = normalize(reference).split()
    right = normalize(hypothesis).split()
    if not left:
        return 0.0 if not right else 1.0
    previous = list(range(len(right) + 1))
    for left_index, left_word in enumerate(left, start=1):
        current = [left_index]
        for right_index, right_word in enumerate(right, start=1):
            current.append(
                min(
                    current[-1] + 1,
                    previous[right_index] + 1,
                    previous[right_index - 1] + (left_word != right_word),
                )
            )
        previous = current
    return previous[-1] / len(left)


def summarize(rows: list[dict]) -> dict:
    latencies = sorted(row["latencyMs"] for row in rows if row["latencyMs"] is not None)
    unknown = [row for row in rows if row["expectedPageId"] is None]
    return {
        "runs": len(rows),
        "successfulRuns": sum(row["error"] is None for row in rows),
        "failures": sum(row["error"] is not None for row in rows),
        "exactMatchRate": sum(
            normalize(row["expectedTranscript"]) == normalize(row["transcript"])
            for row in rows
        ) / len(rows),
        "averageWer": statistics.mean(
            word_error_rate(row["expectedTranscript"], row["transcript"]) for row in rows
        ),
        "intentAccuracy": sum(
            row["expectedPageId"] == row["detectedPageId"] for row in rows
        ) / len(rows),
        "unknownFalsePositiveRate": (
            sum(row["detectedPageId"] is not None for row in unknown) / len(unknown)
            if unknown
            else 0
        ),
        "medianLatencyMs": statistics.median(latencies) if latencies else None,
        "p95LatencyMs": latencies[math.ceil(len(latencies) * 0.95) - 1]
        if latencies
        else None,
    }


def set_run(run, *, size: float = 10, bold: bool = False, color: str = INK, italic: bool = False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_bullet(doc: Document, text: str, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        set_run(paragraph.add_run(bold_prefix), bold=True)
        set_run(paragraph.add_run(text[len(bold_prefix) :]))
    else:
        set_run(paragraph.add_run(text))


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def shade(cell, fill: str):
    shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        cell._tc.get_or_add_tcPr().append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), BORDER)


def set_cell(cell, text: str, *, size: float = 8, bold: bool = False, center: bool = False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(1.5)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0
    set_run(paragraph.add_run(text), size=size, bold=bold)
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), "70")
        node.set(qn("w:type"), "dxa")
    border(cell)


def set_table_geometry(table, widths_inches: list[float]):
    widths = [round(width * 1440) for width in widths_inches]
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        prevent_split(row)
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def audio_duration_seconds(command_id: str) -> float:
    with wave.open(str(AUDIO_DIR / f"{command_id}.wav"), "rb") as wav:
        return wav.getnframes() / wav.getframerate()


def pct(value: float) -> str:
    return f"{value * 100:.1f}%".replace(".", ",")


def ms(value: float | None) -> str:
    return "-" if value is None else f"{value:,.0f} ms".replace(",", ".")


def main():
    data = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    rows = data["observations"]
    assert len(rows) == 90
    assert len({(row["model"], row["commandId"]) for row in rows}) == 90
    summaries = {
        model: summarize([row for row in rows if row["model"] == model])
        for model in MODEL_ORDER
    }
    data["summaries"] = summaries
    RESULTS_PATH.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    document = Document(DOCX_PATH)
    if any(paragraph.text == "7. Benchmark thực nghiệm trên 30 câu TTS" for paragraph in document.paragraphs):
        raise RuntimeError("Benchmark section already exists; refusing to append twice.")

    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)

    document.add_heading("7. Benchmark thực nghiệm trên 30 câu TTS", level=1)
    lead = document.add_paragraph()
    lead.paragraph_format.space_after = Pt(7)
    set_run(
        lead.add_run(
            "Cả ba model được chạy tuần tự trên cùng 30 file WAV mono PCM16, 16 kHz, "
            "giọng Microsoft David (en-US), tổng thời lượng 69,42 giây. Không có audio nào được upload."
        ),
        size=10.5,
    )
    add_bullet(document, "Corpus: 26 navigation commands bao phủ 8 page và 4 câu Unknown.", "Corpus:")
    add_bullet(document, "Moonshine: nhận chunk 100 ms theo thời gian thực, sau đó nhận silence để hoàn tất line.", "Moonshine:")
    add_bullet(document, "Whisper: nhận toàn bộ PCM sau khi file audio đã sẵn sàng và xử lý theo batch.", "Whisper:")
    add_bullet(document, "Không loại bỏ cold first inference; model load time không nằm trong latency từng command.", "Không loại bỏ cold first inference;")

    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    set_run(caption.add_run("Bảng 3. "), bold=True, color=BLUE)
    set_run(caption.add_run("Kết quả tổng hợp từ 30 command trên mỗi model."), color=MUTED)

    headers = ["Model", "Exact / WER", "Intent đúng", "Unknown FP", "Median", "P95", "Lỗi"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, size=8.2, bold=True, center=True)
        shade(table.rows[0].cells[index], HEADER_FILL)
    set_repeat_header(table.rows[0])
    for model in MODEL_ORDER:
        summary = summaries[model]
        row = table.add_row()
        values = [
            MODEL_NAMES[model],
            f"{pct(summary['exactMatchRate'])} / {pct(summary['averageWer'])}",
            pct(summary["intentAccuracy"]),
            pct(summary["unknownFalsePositiveRate"]),
            ms(summary["medianLatencyMs"]),
            ms(summary["p95LatencyMs"]),
            str(summary["failures"]),
        ]
        for index, value in enumerate(values):
            set_cell(row.cells[index], value, size=8.1, bold=index == 0, center=index > 0)
            if model == "small-streaming":
                shade(row.cells[index], SMALL_FILL)
    set_table_geometry(table, [1.35, 1.05, 0.82, 0.78, 0.9, 0.9, 0.7])

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(8)
    set_run(note.add_run("Định nghĩa latency: "), bold=True, color=BLUE, size=9)
    set_run(
        note.add_run(
            "Moonshine tính từ chunk đầu tiên đến final transcript nên bao gồm thời lượng phát audio; "
            "Whisper tính từ lúc nhận toàn bộ buffer đến result. Vì cơ chế input khác nhau, số liệu thể hiện "
            "pipeline thực tế của project nhưng không phải benchmark kernel thuần túy."
        ),
        size=9,
        italic=True,
        color=MUTED,
    )

    document.add_heading("8. Diễn giải kết quả", level=1)
    add_bullet(document, "ASR trên TTS sạch: cả ba đạt 96,7% exact match và WER trung bình 1,7%; không có failure.", "ASR trên TTS sạch:")
    add_bullet(document, "Intent: cả ba đạt 28/30 (93,3%). Hai miss giống nhau đến từ local matcher, không phải khác biệt model.", "Intent:")
    add_bullet(document, "Câu “reports dashboard” bị trả Unknown do xung đột alias Dashboard/Analytics; câu “Open Project integration” không khớp alias “OpenProject integration”.", "Câu “reports dashboard”")
    add_bullet(document, "Latency: Small nhanh nhất (median 2.848 ms), sau đó Medium (3.432 ms), Whisper Base (5.435 ms).", "Latency:")
    add_bullet(document, "Khuyến nghị sau benchmark: giữ Moonshine Small v2 làm mặc định. Medium không tăng accuracy trên corpus sạch này nhưng chậm hơn; Whisper Base chậm nhất và không có partial streaming.", "Khuyến nghị sau benchmark:")

    limitation = document.add_paragraph()
    limitation.paragraph_format.space_before = Pt(7)
    limitation.paragraph_format.space_after = Pt(0)
    set_run(limitation.add_run("Giới hạn: "), bold=True, color=BLUE)
    set_run(
        limitation.add_run(
            "TTS sạch không đại diện cho microphone gain thấp, noise, reverberation, accent hoặc speech tự nhiên. "
            "Kết quả này phù hợp để so sánh pipeline trong điều kiện tái lập; quyết định production vẫn cần một corpus thu từ người dùng thật."
        )
    )

    detail_section = document.add_section(WD_SECTION.NEW_PAGE)
    detail_section.orientation = WD_ORIENT.LANDSCAPE
    detail_section.page_width = Inches(11)
    detail_section.page_height = Inches(8.5)
    detail_section.left_margin = Inches(0.65)
    detail_section.right_margin = Inches(0.65)
    detail_section.top_margin = Inches(0.65)
    detail_section.bottom_margin = Inches(0.65)

    document.add_heading("9. Kết quả chi tiết 30 command", level=1)
    detail_caption = document.add_paragraph()
    detail_caption.paragraph_format.space_after = Pt(5)
    set_run(
        detail_caption.add_run(
            "Mỗi ô model hiển thị transcript, latency và kết quả intent. “OK” nghĩa là page được phát hiện đúng ground truth."
        ),
        size=9,
        color=MUTED,
    )

    detail_headers = ["ID", "Ground truth command", "Page", "Audio", "Small v2", "Medium v2", "Whisper Base.en"]
    detail = document.add_table(rows=1, cols=len(detail_headers))
    detail.style = "Table Grid"
    for index, header in enumerate(detail_headers):
        set_cell(detail.rows[0].cells[index], header, size=7.6, bold=True, center=True)
        shade(detail.rows[0].cells[index], HEADER_FILL)
    set_repeat_header(detail.rows[0])

    by_key = {(row["model"], row["commandId"]): row for row in rows}
    for entry in data["corpus"]:
        row = detail.add_row()
        expected_page = entry["expectedPageId"] or "Unknown"
        set_cell(row.cells[0], entry["id"].replace("cmd-", ""), size=7.2, center=True)
        set_cell(row.cells[1], entry["text"], size=7.2)
        set_cell(row.cells[2], expected_page.title(), size=7.2, center=True)
        set_cell(row.cells[3], f"{audio_duration_seconds(entry['id']):.2f} s", size=7.2, center=True)
        for column, model in enumerate(MODEL_ORDER, start=4):
            result = by_key[(model, entry["id"])]
            intent_ok = result["detectedPageId"] == result["expectedPageId"]
            status = "OK" if intent_ok else f"MISS → {result['detectedPageId'] or 'Unknown'}"
            set_cell(
                row.cells[column],
                f"{result['transcript']}\n{ms(result['latencyMs'])} | {status}",
                size=7.0,
            )
            if model == "small-streaming":
                shade(row.cells[column], SMALL_FILL)
    set_table_geometry(detail, [0.42, 2.05, 0.72, 0.58, 1.98, 1.98, 1.97])

    document.save(DOCX_PATH)
    with ZipFile(DOCX_PATH) as archive:
        assert archive.testzip() is None
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
